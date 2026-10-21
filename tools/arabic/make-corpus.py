# Second half of the Arabic editing corpus (tests/fixtures/arabic/): the same paragraphs through Microsoft Word (COM,
# "Save as PDF"), LibreOffice headless (the Office add-on in the app data folder, or an installed LibreOffice) and
# MuPDF's own HTML layout (PyMuPDF Story). Run with the scratch venv:  <venv>/Scripts/python tools/arabic/make-corpus.py
# Needs: pymupdf (system Python has 1.28), python-docx, pywin32 (Word only). Texts come from texts.mjs via node.
import json, os, subprocess, sys, tempfile, pathlib, shutil

ROOT = pathlib.Path(__file__).resolve().parents[2]
OUT = ROOT / 'tests' / 'fixtures' / 'arabic'
# Microsoft Store Python sandboxes %APPDATA% and %TEMP%: pass WARAQ_SOFFICE and ARABIC_TMP from the shell when needed
SOFFICE = [p for p in [pathlib.Path(os.environ.get('WARAQ_SOFFICE', 'x')), pathlib.Path(os.environ['APPDATA']) / 'com.waraq.app/engines/office/libreoffice/program/soffice.exe',
                       pathlib.Path('C:/Program Files/LibreOffice/program/soffice.exe')] if p.exists()]
TMP = pathlib.Path(os.environ.get('ARABIC_TMP') or pathlib.Path(tempfile.gettempdir()) / 'waraq-arabic-corpus')
TMP.mkdir(exist_ok=True)
TEXTS = json.loads(subprocess.check_output(['node', '--input-type=module', '-e',
    "import {TEXTS} from './tests/fixtures/arabic/texts.mjs'; console.log(JSON.stringify(TEXTS))"], cwd=ROOT, encoding='utf-8'))
LINES = [dict(t, text='    '.join(t['cells'])) if 'cells' in t else t for t in TEXTS]
slug = lambda s: ''.join(c if c.isalnum() else '-' for c in s.lower()).strip('-')
def save(name, src):
    shutil.copyfile(src, OUT / name); print(name, round(os.path.getsize(src) / 1024), 'kB')

# ---- LibreOffice: flat ODT (no library needed) -> PDF ----
LO_FONTS = ['Traditional Arabic', 'Arial', 'Times New Roman', 'Sakkal Majalla', 'Tahoma', 'Simplified Arabic']
def fodt(font):
    esc = lambda s: s.replace('&', '&amp;').replace('<', '&lt;')
    styles, paras = [], []
    for i, t in enumerate(LINES):
        size = t.get('size', 13)
        props = f'fo:font-size="{size}pt" style:font-size-complex="{size}pt" style:font-name="{font}" style:font-name-complex="{font}"'
        if t.get('bold'): props += ' fo:font-weight="bold" style:font-weight-complex="bold"'
        if t.get('italic'): props += ' fo:font-style="italic" style:font-style-complex="italic"'
        if t.get('color'): props += f' fo:color="{t["color"]}"'
        align = 'justify' if t['kind'] == 'justified' else 'end'
        styles.append(f'<style:style style:name="P{i}" style:family="paragraph"><style:paragraph-properties fo:text-align="{align}" style:writing-mode="rl-tb" fo:margin-bottom="4pt"/><style:text-properties {props}/></style:style>')
        paras.append(f'<text:p text:style-name="P{i}">{esc(t["text"])}</text:p>')
    return f'''<?xml version="1.0" encoding="UTF-8"?>
<office:document xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" office:version="1.2" office:mimetype="application/vnd.oasis.opendocument.text">
<office:font-face-decls><style:font-face style:name="{font}" svg:font-family="'{font}'"/></office:font-face-decls>
<office:automatic-styles>{''.join(styles)}</office:automatic-styles>
<office:body><office:text>{''.join(paras)}</office:text></office:body></office:document>'''

def libreoffice():
    for font in LO_FONTS: (TMP / f'lo-{slug(font)}.fodt').write_text(fodt(font), encoding='utf-8')
    if not SOFFICE: print('LibreOffice not found: .fodt files written to', TMP, '- convert them with soffice --headless --convert-to pdf'); return
    profile = (TMP / 'lo-profile').as_uri()
    for font in LO_FONTS:
        src = TMP / f'lo-{slug(font)}.fodt'
        subprocess.run([str(SOFFICE[0]), '--headless', '--norestore', f'-env:UserInstallation={profile}', '--convert-to', 'pdf', '--outdir', str(TMP), str(src)],
                       capture_output=True, timeout=180)
        pdf = src.with_suffix('.pdf')
        if pdf.exists(): save(f'libreoffice-{slug(font)}.pdf', pdf)
        else: print('LibreOffice failed for', font)

# ---- Microsoft Word: .docx (python-docx) -> Word COM "Save as PDF" ----
WORD_FONTS = (os.environ.get('WORD_FONTS') or 'Traditional Arabic,Arial,Times New Roman,Sakkal Majalla,Calibri,Simplified Arabic,Dubai').split(',')
def docx_file(font, path):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt, RGBColor
    d = docx.Document()
    for t in LINES:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if t['kind'] == 'justified' else WD_ALIGN_PARAGRAPH.RIGHT
        bidi = OxmlElement('w:bidi'); p._p.get_or_add_pPr().append(bidi)  # right-to-left paragraph
        r = p.add_run(t['text']); r.font.size = Pt(t.get('size', 13)); r.font.name = font
        rpr = r._r.get_or_add_rPr()
        fonts = rpr.find(qn('w:rFonts')); fonts = OxmlElement('w:rFonts') if fonts is None else fonts; fonts.set(qn('w:cs'), font); fonts.set(qn('w:ascii'), font); fonts.set(qn('w:hAnsi'), font); rpr.append(fonts)
        for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'): fonts.attrib.pop(qn(a), None)  # theme fonts override named ones in Word
        fonts.set(qn('w:hint'), 'cs')
        rtl = OxmlElement('w:rtl'); rpr.append(rtl)
        szcs = OxmlElement('w:szCs'); szcs.set(qn('w:val'), str(t.get('size', 13) * 2)); rpr.append(szcs)
        if t.get('bold'): r.bold = True; bcs = OxmlElement('w:bCs'); rpr.append(bcs)
        if t.get('italic'): r.italic = True; ics = OxmlElement('w:iCs'); rpr.append(ics)
        if t.get('color'): c = t['color'].lstrip('#'); r.font.color.rgb = RGBColor(int(c[:2], 16), int(c[2:4], 16), int(c[4:], 16))
    d.save(path)

def word():
    try:
        import win32com.client
        app = win32com.client.Dispatch('Word.Application')
    except Exception as e:
        print('Word not available:', e); return
    app.Visible = False
    try:
        for font in WORD_FONTS:
            src = TMP / f'word-{slug(font)}.docx'; docx_file(font, src)
            pdf = src.with_suffix('.pdf')
            doc = app.Documents.Open(str(src), ReadOnly=True)
            doc.ExportAsFixedFormat(str(pdf), 17)  # wdExportFormatPDF
            doc.Close(False)
            save(f'word-{slug(font)}.pdf', pdf)
    finally:
        app.Quit()

# ---- MuPDF's own layout (Story): the Edge HTML, fonts by file ----
def story():
    import fitz
    fonts_dir = pathlib.Path(os.environ.get('ARABIC_FONTS') or pathlib.Path(os.environ['LOCALAPPDATA']) / 'Temp/claude/C--Users-Abdullah-Documents-ChatGPT-PDF-Open-Source-waraq-next/d5d66cc1-abe4-4e52-abe7-f46d0d5ffe77/scratchpad/fonts')
    picks = [('Noto Naskh Arabic', ROOT / 'public/fonts/NotoNaskhArabic-Regular.ttf'), ('Amiri', fonts_dir / 'Amiri-Regular.ttf'), ('Tajawal', ROOT / 'public/fonts/Tajawal-Regular.ttf')]
    for family, file in picks:
        if not file.exists(): continue
        html = subprocess.check_output(['node', '--input-type=module', '-e',
            f"import {{html}} from './tools/arabic/make-corpus.mjs'; console.log(html({family!r}))"], cwd=ROOT, encoding='utf-8')
        arch = fitz.Archive(str(file.parent))
        css = f'@font-face {{ font-family: "{family}"; src: url("{file.name}"); }}'
        st = fitz.Story(html=html, user_css=css, archive=arch)
        pdf = TMP / f'mupdf-story-{slug(family)}.pdf'
        writer = fitz.DocumentWriter(str(pdf))
        more = True
        while more:
            dev = writer.begin_page(fitz.paper_rect('a4'))
            more, _ = st.place(fitz.paper_rect('a4') + (50, 50, -50, -50)); st.draw(dev)
            writer.end_page()
        writer.close()
        save(f'mupdf-story-{slug(family)}.pdf', pdf)

if __name__ == '__main__':
    what = sys.argv[1] if len(sys.argv) > 1 else 'all'
    if what in ('all', 'libreoffice'): libreoffice()
    if what in ('all', 'word'): word()
    if what in ('all', 'story'): story()
